"""Export exact Beamer pages to a portable PowerPoint; source stays editable.

Run with the bundled document Python (python-pptx and pypdf required). Each
slide is one high-resolution rendered page, with searchable source text in
notes. This deliberately preserves the supplied Beamer design; the .tex bundle
is the editable master and the .pptx is a presentation-compatible rendition.
"""
import hashlib,json,subprocess,zipfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
FOLDER=ROOT/'presentations/math_tokenizer'


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    output=FOLDER/'output';output.mkdir(exist_ok=True)
    name='Design_of_a_Specialized_Tokenizer_for_Mathematical_Text'
    pdf=output/(name+'.pdf');pdf.write_bytes((FOLDER/'build/mathematical_tokenizer.pdf').read_bytes())
    pages=FOLDER/'qa/pptx_pages';pages.mkdir(exist_ok=True)
    subprocess.run(['pdftoppm','-png','-scale-to','2000',str(pdf),str(pages/'slide')],check=True,stdout=subprocess.DEVNULL)
    images=sorted(pages.glob('slide-*.png'));reader=PdfReader(pdf)
    assert len(images)==len(reader.pages)==43
    presentation=Presentation();presentation.slide_width=Inches(10);presentation.slide_height=Inches(7.5)
    presentation.core_properties.title='Design of a Specialized Tokenizer for Mathematical Text'
    presentation.core_properties.author='Tongren Xiao'
    presentation.core_properties.subject='Mathematical tokenizer research; exact Beamer rendition'
    for image,page in zip(images,reader.pages):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image),0,0,width=presentation.slide_width,height=presentation.slide_height)
        slide.notes_slide.notes_text_frame.text=(page.extract_text() or '')+'\n\nEditable master: mathematical_tokenizer.tex in the companion Beamer source bundle. The PowerPoint slide preserves a rendered Beamer page.'
    pptx=output/(name+'.pptx');presentation.save(pptx)
    reloaded=Presentation(pptx);assert len(reloaded.slides)==len(images)
    for slide,image in zip(reloaded.slides,images):
        assert len(slide.shapes)==1 and slide.shapes[0].image.blob==image.read_bytes()
        assert slide.shapes[0].left==slide.shapes[0].top==0
    (FOLDER/'source/BUILD.md').write_text('# Editable Beamer master\n\nRun `latexmk -xelatex mathematical_tokenizer.tex` from this folder. Requires XeLaTeX, Beamer and ctex. The source and template remain editable; assets are vector PDF data figures. The cover information and October 10, 2026 date are preserved from the reference at user request; the research was updated October 13, 2026.\n\nThe accompanying PowerPoint contains exact high-resolution renderings, with source text in notes. Edit the TeX master to change content.\n')
    bundle=output/'Mathematical_Tokenizer_Beamer_Source.zip'
    with zipfile.ZipFile(bundle,'w',compression=zipfile.ZIP_DEFLATED) as archive:
        for p in sorted((FOLDER/'source').rglob('*')):
            if p.is_file():archive.write(p,p.relative_to(FOLDER/'source'))
    # Do not ship an apparently standalone .tex file without its dependencies.
    old=output/(name+'.tex')
    if old.exists():old.unlink()
    record={'slides':43,'aspect_ratio':'4:3','pptx_is_exact_page_rendition':True,'editable_master':'source/mathematical_tokenizer.tex',
        'slide_image_bytes_verified':True,'pdf_sha256':sha(pdf),'pptx_sha256':sha(pptx),'source_bundle_sha256':sha(bundle)}
    (FOLDER/'qa/export_verification.json').write_text(json.dumps(record,indent=2)+'\n')
    print('Exported PDF, 43-slide PowerPoint and editable Beamer bundle',flush=True)


if __name__=='__main__':main()
